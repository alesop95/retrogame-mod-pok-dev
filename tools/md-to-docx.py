#!/usr/bin/env python3
"""Converte un file Markdown del progetto in un .docx parallelo, con titoli, prosa, citazioni, tabelle, blocchi di codice e figure.

Perche' esiste
--------------

Alcuni documenti del progetto servono anche fuori dall'editor, stampati o aperti su un telefono, e il proprietario ha chiesto il 2026-09-23 una copia .docx della guida al Parco Lotta tenuta allineata man mano. Una copia fatta a mano diverge alla prima modifica, quindi la copia si genera: il Markdown resta la fonte e il .docx e' un derivato, che non si modifica e non si legge per diff. Sulla macchina non c'e' pandoc, e il sottoinsieme di Markdown che il progetto usa e' piccolo e vincolato dalla regola di stile, quindi basta un convertitore di poche centinaia di righe sulla libreria python-docx.

Che cosa copre, e che cosa no
------------------------------

Copre cio' che la regola `interaction-style.md` ammette: titoli da uno a quattro livelli, paragrafi su riga unica, citazioni con il prefisso `>`, tabelle con la riga di separazione, blocchi recintati, immagini nella forma `![didascalia](percorso)` su una riga propria, e dentro il testo il codice fra apici inversi, il corsivo con l'asterisco singolo e il grassetto con il doppio. I commenti HTML, cioe' i marcatori dei blocchi generati, si saltano. Non copre elenchi annidati complessi, note a pie' di pagina, collegamenti come oggetti cliccabili: il testo del collegamento resta, l'indirizzo si perde. Un costrutto non riconosciuto diventa un paragrafo di testo, mai un errore silenzioso che lo fa sparire.

Uso
---

    python tools/md-to-docx.py GUIDA.md
    python tools/md-to-docx.py GUIDA.md --out GUIDA.docx
"""

import argparse
import datetime
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")


def testo_inline(paragrafo, testo, base_bold=False):
    for pezzo in INLINE.split(testo):
        if not pezzo:
            continue
        if pezzo.startswith("`") and pezzo.endswith("`"):
            run = paragrafo.add_run(pezzo[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif pezzo.startswith("**") and pezzo.endswith("**"):
            run = paragrafo.add_run(pezzo[2:-2])
            run.bold = True
        elif pezzo.startswith("*") and pezzo.endswith("*") and len(pezzo) > 2:
            run = paragrafo.add_run(pezzo[1:-1])
            run.italic = True
        elif pezzo.startswith("["):
            run = paragrafo.add_run(re.match(r"\[([^\]]+)\]", pezzo).group(1))
        else:
            run = paragrafo.add_run(pezzo)
        if base_bold:
            run.bold = True


def celle(riga):
    return [c.strip() for c in riga.strip().strip("|").split("|")]


def converti(sorgente, destinazione):
    righe = Path(sorgente).read_text(encoding="utf-8").split("\n")
    base = Path(sorgente).parent
    doc = Document()
    stile = doc.styles["Normal"]
    stile.font.name = "Calibri"
    stile.font.size = Pt(10.5)
    for sezione in doc.sections:
        sezione.left_margin = sezione.right_margin = Cm(2)
        sezione.top_margin = sezione.bottom_margin = Cm(1.8)
    i = 0
    while i < len(righe):
        r = righe[i]
        s = r.strip()
        if not s:
            i += 1
            continue
        if s.startswith("<!--"):
            while "-->" not in righe[i]:
                i += 1
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            blocco = []
            while i < len(righe) and not righe[i].strip().startswith("```"):
                blocco.append(righe[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(blocco))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)) - 1, 3) if len(m.group(1)) > 1 else 0)
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", s)
        if m:
            percorso = base.joinpath(m.group(2))
            if percorso.exists():
                doc.add_picture(str(percorso), width=Cm(17))
                didascalia = doc.add_paragraph()
                run = didascalia.add_run(m.group(1))
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
            else:
                doc.add_paragraph("[figura mancante: %s]" % m.group(2))
            i += 1
            continue
        if s.startswith("|"):
            tabella = []
            while i < len(righe) and righe[i].strip().startswith("|"):
                if not re.match(r"^\|\s*-+", righe[i].strip()):
                    tabella.append(celle(righe[i]))
                i += 1
            colonne = max(len(t) for t in tabella)
            t = doc.add_table(rows=len(tabella), cols=colonne)
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, riga in enumerate(tabella):
                for ci in range(colonne):
                    cella = t.cell(ri, ci)
                    cella.paragraphs[0].text = ""
                    testo_inline(cella.paragraphs[0], riga[ci] if ci < len(riga) else "", base_bold=(ri == 0))
                    for run in cella.paragraphs[0].runs:
                        run.font.size = Pt(8.5)
            doc.add_paragraph()
            continue
        if s.startswith(">"):
            contenuto = s.lstrip(">").strip()
            if contenuto:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                testo_inline(p, contenuto)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
            i += 1
            continue
        p = doc.add_paragraph()
        testo_inline(p, s)
        i += 1
    # Il file deve restare identico fra due corse sullo stesso sorgente, altrimenti ogni rigenerazione
    # e' un binario nuovo nel repository anche quando nulla e' cambiato. python-docx scrive la data
    # corrente nelle proprieta' del documento e l'ora di ogni voce dell'archivio: si fissano entrambe.
    fissa = datetime.datetime(2000, 1, 1)
    doc.core_properties.created = fissa
    doc.core_properties.modified = fissa
    doc.core_properties.last_modified_by = ""
    doc.core_properties.author = ""
    memoria = io.BytesIO()
    doc.save(memoria)
    memoria.seek(0)
    uscita = io.BytesIO()
    with zipfile.ZipFile(memoria) as zin, zipfile.ZipFile(uscita, "w", zipfile.ZIP_DEFLATED) as zout:
        for nome in sorted(zin.namelist()):
            voce = zipfile.ZipInfo(nome, date_time=(2000, 1, 1, 0, 0, 0))
            voce.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(voce, zin.read(nome))
    Path(destinazione).write_bytes(uscita.getvalue())


def main():
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("sorgente")
    p.add_argument("--out")
    args = p.parse_args()
    destinazione = Path(args.out) if args.out else Path(args.sorgente).with_suffix(".docx")
    converti(args.sorgente, destinazione)
    print("scritto %s" % destinazione)


if __name__ == "__main__":
    main()
