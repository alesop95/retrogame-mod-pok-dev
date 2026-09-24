#!/usr/bin/env python3
"""Genera la mappa stampabile dei quattordici box di Smeraldo: che cosa c'e' in ogni posizione, con l'illustrazione di Sugimori, e da dove viene.

Perche' esiste
--------------

Dopo ADR-076, ADR-077 e ADR-078 il deposito e' pieno, 420 su 420, e mescola sette provenienze: la collezione di catture vere, gli eventi in tutte le lingue, gli incontri da biglietto, gli scambi in gioco, il Seedot e il Lotad piu' grandi possibili, il Wynaut dell'Isola Miraggio e il lotto del Parco Lotta. Chi apre il PC del gioco vede soltanto trenta icone per box e non sa quale di tre Latios sia l'evento e quale la cattura. Questo strumento risponde posizione per posizione, e lo fa leggendo i byte del salvataggio invece di fidarsi della regola che li ha disposti: la provenienza di ogni posizione si riconosce confrontando il record con i file dei lotti, e i dati leggibili vengono dal dump di PKHeX.

Il dump puo' appartenere a un altro file, purche' contenga gli stessi record, perche' lo strumento lo accoppia per contenuto e non per posizione; se ne possono passare piu' d'uno. Un record che nessun dump contiene, come un esemplare appena aggiunto e non ancora giudicato, si descrive dai suoi byte e dalle tabelle del progetto, e lo strumento lo dichiara a terminale.

Per gli eventi la tabella riporta la storia della distribuzione, cioe' quando, dove e come fu consegnata e perche' e' notevole, da `recreate-pokemon-distributions-events/provenienze-eventi.json`, che e' la stessa fonte autorata del catalogo degli eventi, con l'evento individuato dalla coppia allenatore e identificativo della tabella del verificatore. Per non ripetere lo stesso testo, la storia si scrive alla prima posizione del box in cui l'evento compare, e le successive vi rimandano.

La stampa
---------

Il proprietario ha fissato il formato il 2026-09-23: foglio orizzontale, per ogni box una pagina con la figura e una con la tabella, mai di piu'; la legenda dei colori e il riepilogo una volta sola in testa; carattere sempre nero; nessuna colonna di legalita', perche' il deposito e' tutto legale e la colonna sarebbe costante. La copia `.docx` si costruisce qui con python-docx invece che dal convertitore generico, perche' servono l'orientamento orizzontale e il salto di pagina per box; poi Word la esporta in PDF e lo strumento conta le pagine di ogni tabella, e dove una tabella sfora rimpicciolisce il carattere di quel box e ricomincia. Le figure usano Arial e le tabelle Arial Narrow, caratteri da ufficio invece di quello predefinito di matplotlib.

Le illustrazioni vengono dalla raccolta locale del proprietario e restano fuori da git con le figure e le copie da stampare, per ADR-005; il documento Markdown generato si versiona.

Uso
---

    python gba-save-extraction-smeraldo/tools/emerald_mappa_box.py SALVATAGGIO.sav --dump DUMP.csv FILE_DEL_DUMP.sav [--dump ...]

Scrive `gba-save-extraction-smeraldo/MAPPA-BOX-SMERALDO.md` e `MAPPA-BOX-RUBINO.md`, le figure `figure/mappa-box-NN.png` e `figure/mappa-rubino-box-NN.png`, e la copia da stampare della collezione intera, `MAPPA-BOX-COLLEZIONE.docx` e `MAPPA-BOX-COLLEZIONE.pdf`: prima Smeraldo, poi il Rubino. Dal 2026-09-24 il Rubino entra nella stampa con la disposizione prevista del complemento di ADR-080, presa da `disposizione_rubino`, finche' la cartuccia non e' scritta.

La disposizione del Rubino segue i gruppi della sua legenda, cioè statici, esclusivi, specie, portatori di mosse e fiocchi, Unown, Colosseum, XD ed evento, e dentro ciascun gruppo il numero del Pokédex. È la disposizione che la scrittura dovrà usare.
"""

import argparse
import collections
import csv
import importlib.util
import json
import os
import re
import subprocess
import textwrap
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402
from PIL import Image  # noqa: E402

# Arial per le figure; il font giapponese di Windows fa da riserva per gli allenatori in kana.
plt.rcParams["font.family"] = ["Arial", "Yu Gothic", "MS Gothic"]

CARTELLA = Path(__file__).resolve().parents[1]
RADICE = CARTELLA.parent
sys.path.insert(0, str(RADICE.joinpath("pokemon-gen12-gen3-bridge-original-hardware")))

from pokebridge import charmap, gen3, save3  # noqa: E402

NOTE = RADICE.joinpath("_notes")
FIGURE = CARTELLA.joinpath("figure")
USCITA = CARTELLA.joinpath("MAPPA-BOX-SMERALDO.md")
PKHEX = NOTE.joinpath("fonti", "pkhex")
_SUGIMORI_RELATIVO = Path("My files", "Sugimori Pokémon Gen1-9 DLC3 Organized",
                          "Sugimori Pokémon Gen1-9 DLC3 Organized", "Pokémon By Generation")
_SUGIMORI_TROVATI = sorted(Path.home().joinpath("Proton Drive").glob("*/" + _SUGIMORI_RELATIVO.as_posix()))
SUGIMORI = Path(os.environ["SUGIMORI_ROOT"]) if "SUGIMORI_ROOT" in os.environ else (
    _SUGIMORI_TROVATI[0] if len(_SUGIMORI_TROVATI) == 1 else Path.home().joinpath("Proton Drive", _SUGIMORI_RELATIVO))
PER_BOX = 30
COLONNE = 6
VUOTO = bytes(save3.RECORD)
# Nomi di gioco gia' presenti nella mappa pubblica; ogni nuovo OT della collezione
# richiede una verifica prima di entrare nella mappa, per non esporre persone terze.
OT_COLLEZIONE_PUBBLICI = {
    "ALEX", "ALESSIO", "AXEL", "CICCIO", "10ANNI", "Al Paci",
    "DANIELE", "DONTAE", "MARCO", "MATT", "Gian",
}

# Sette provenienze, sette slot categoriali nell'ordine fisso della tavolozza di riferimento della
# skill di visualizzazione. Il testo e' sempre nero, per scelta del proprietario, quindi le caselle
# usano una tinta schiarita di ciascun colore, sulla quale il nero si legge; il colore non e' mai il
# solo segnale, perche' ogni casella porta la provenienza scritta.
PROVENIENZE = [
    ("collezione", "Collezione", "Collezione", "#2a78d6"),
    ("evento", "Evento", "Evento", "#eb6834"),
    ("biglietto", "Incontro da biglietto", "Incontro da biglietto", "#1baf7a"),
    ("scambio", "Scambio in gioco", "Scambio in gioco", "#eda100"),
    ("gigante", "Seedot e Lotad più grandi possibili", "Il più grande possibile", "#e87ba4"),
    ("selvatico", "Wynaut dell'Isola Miraggio", "Isola Miraggio", "#4a3aa7"),
    ("lotto", "Parco Lotta", "Parco Lotta", "#008300"),
]
# Il Rubino di prova riceve il complemento di ADR-080, e le sue provenienze sono altre; hanno una legenda
# propria, quindi riusano gli otto colori della tavolozza nello stesso ordine fisso.
PROVENIENZE_RUBINO = [
    ("r-statico", "Incontri statici, doni, uova e incontri speciali", "Statico, dono o speciale", "#2a78d6"),
    ("r-esclusivo", "Esclusivi di versione", "Esclusivo", "#eb6834"),
    ("r-specie", "Specie che completano il Pokédex", "Pokédex", "#1baf7a"),
    ("r-portatore", "Mosse perdute e fiocchi", "Mossa o fiocchi", "#eda100"),
    ("r-unown", "Forme di Unown", "Unown", "#e87ba4"),
    ("r-colosseum", "Colosseum", "Colosseum", "#4a3aa7"),
    ("r-xd", "XD", "XD", "#008300"),
    ("r-evento", "Evento", "Evento", "#e34948"),
]
ETICHETTA = {k: e for k, e, _, _ in PROVENIENZE + PROVENIENZE_RUBINO}
BANDA = {k: b for k, _, b, _ in PROVENIENZE + PROVENIENZE_RUBINO}
LETTERE_UNOWN = "ABCDEFGHIJKLMNOPQRSTUVWXYZ!?"
RUBINO = RADICE.joinpath("_notes", "lotto-complemento-rubino", "esemplari")
# Gli sfondi di Rubino sono sedici, nello stesso ordine di Smeraldo, e lo sfondo Amici non c'e': e' un'aggiunta
# di Smeraldo, legata alla frase di Walda, e `gWallpaperTable` in `src/pokemon_storage_system_4.c` di pokeruby
# elenca soltanto Forest, City, Desert, Savanna, Crag, Volcano, Snow, Cave, Beach, Seafloor, River, Sky,
# Polkadot, Pokecenter, Machine e Plain. Ogni box prende il primo sfondo libero fra quelli pensati per il
# gruppo che vi prevale: citta' e centro per gli statici, savana per gli esclusivi, i paesaggi per il
# Pokedex, il pois per gli Unown, il deserto e il vulcano per Colosseum, che si gioca nel deserto di Orre,
# la macchina e il fondale per XD, il cielo per l'evento.
NOMI_SFONDI_RUBINO = ["Forest", "City", "Desert", "Savanna", "Crag", "Volcano", "Snow", "Cave", "Beach", "Seafloor", "River",
                      "Sky", "Polkadot", "Pokecenter", "Machine", "Plain"]
SFONDI_PER_GRUPPO = {"r-statico": [1, 13], "r-esclusivo": [3, 4], "r-specie": [0, 10, 8, 6], "r-portatore": [15, 7],
                     "r-unown": [12, 4], "r-colosseum": [2, 5], "r-xd": [14, 9, 7], "r-evento": [11]}


def sfondi_rubino(per_box):
    """Lo sfondo di ciascun box del Rubino, dal gruppo che vi prevale, senza ripetere uno sfondo gia' usato."""
    usati, fuori = set(), []
    for contenuto in per_box:
        conti = collections.Counter(v["provenienza"] for v in contenuto if v)
        # prima gli sfondi del gruppo che prevale, poi quelli degli altri gruppi presenti nel box
        candidati = [x for g, _ in conti.most_common() for x in SFONDI_PER_GRUPPO[g]]
        scelta = next((x for x in candidati if x not in usati), None)
        if scelta is None:
            scelta = next(x for x in range(16) if x not in usati)
        usati.add(scelta)
        fuori.append(scelta)
    return fuori
USCITA_RUBINO = CARTELLA.joinpath("MAPPA-BOX-RUBINO.md")
STAMPA = CARTELLA.joinpath("MAPPA-BOX-COLLEZIONE")
NERO = "#000000"
SUPERFICIE = "#ffffff"
BORDO = "#bdbcb6"
FORME_DEOXYS = {"1": "Attack", "2": "Defense", "3": "Speed"}
# La raccolta di Sugimori ha una sola illustrazione di Unown, la F, e con quella tutte le 28 caselle sembravano
# la stessa lettera. Per Unown si usa l'artwork per forma che PKHeX porta con se', un file per lettera: la A e'
# `a_201.png`, le altre `a_201-N.png` con N la forma, da 1 per la B a 27 per il punto interrogativo.
ARTWORK_PKHEX = RADICE.joinpath("_notes", "fonti", "pkhex", "PKHeX.Drawing.PokeSprite", "Resources", "img", "Artwork Pokemon Sprites")
# Le note delle provenienze mescolano la storia della distribuzione con osservazioni sul lavoro del
# progetto, cioe' su tabelle, verificatore e fonti. La stampa e' per chi guarda la collezione e ne
# tiene soltanto la storia: una frase che nomina uno di questi termini resta nel catalogo degli eventi
# e non entra qui. Per la stessa ragione non entrano le divergenze fra le fonti.
DI_LAVORO = re.compile(r"verificator|progett|tabella|PKHeX|enciclopedi|fonte|fonti|strumento|campo", re.IGNORECASE)


def tinta(esadecimale, quota=0.42):
    """Il colore mescolato al bianco: quota e' la parte di colore che resta."""
    r, g, b = (int(esadecimale[i:i + 2], 16) for i in (1, 3, 5))
    return "#%02x%02x%02x" % tuple(round(255 - (255 - c) * quota) for c in (r, g, b))


TINTA = {k: tinta(c) for k, _, _, c in PROVENIENZE + PROVENIENZE_RUBINO}


def _modulo(percorso, nome):
    spec = importlib.util.spec_from_file_location(nome, percorso)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


def storie_degli_eventi():
    """{record cifrato: chiave dell'evento} e le provenienze storiche per chiave.

    Il numero nel nome del file non e' sempre l'indice della tabella di PKHeX: il generatore ha
    numerato a parte, con quattro cifre, le quattro uova di Pokemon Box, che nella tabella sono le voci
    123-126, e i file da 123 in poi sono quindi spostati di quattro. L'abbinamento non si fida percio'
    del numero ma lo verifica sul contenuto: una voce e' candidata solo se specie, livello e mosse del
    file coincidono con le sue, e fra le candidate si sceglie quella al numero del file, oppure a quel
    numero piu' quattro; per le quattro uova, l'unica candidata dell'allenatore del blocco di Pokemon Box.
    Se nessuna regola da' una sola voce, lo strumento si ferma invece di attribuire una storia a caso.
    """
    catalogo = _modulo(RADICE.joinpath("tools", "catalogo-eventi.py"), "catalogo_eventi")
    generatore = catalogo.carica_generatore()
    voci = generatore.voci_wc3(str(PKHEX))
    interno = {v: k for k, v in generatore.nazionale_verso_interno(str(NOTE.joinpath("fonti", "ace-builder"))).items()}
    provenienze = json.loads(Path(catalogo.PROVENIENZE).read_text(encoding="utf-8"))
    per_record = {}
    for f in sorted(NOTE.joinpath("lotto-eventi").glob("*.pk3")):
        numero = int(f.stem.split("-")[0])
        cifrato = NOTE.joinpath("lotto-eventi", "forma-cifrata", f.stem + ".ek3").read_bytes()
        mon = gen3.Gen3Mon.from_bytes(cifrato)
        mosse = sorted(m for m in mon.attacks.moves if m)
        candidate = [i for i, v in enumerate(voci) if v["nazionale"] == interno[mon.growth.species]
                     and sorted(m for m in v.get("mosse", []) if m) == mosse]
        scelte = [i for i in (numero, numero + 4) if i in candidate][:1] or                  [i for i in candidate if "ＡＺＵＳＡ" in voci[i]["ot"]]
        if len(scelte) != 1:
            sys.exit("il file %s non si abbina a una sola voce della tabella: candidate %s" % (f.name, candidate))
        per_record[cifrato] = (catalogo.chiave(voci[scelte[0]]), voci[scelte[0]]["commento"])
    return per_record, provenienze["gruppi"]


def provenienza_evento(chiave, commento, provenienze):
    """La voce storica dell'evento e la sua identita' per il raggruppamento nel box: la chiave, oppure il sottogruppo quando la chiave ne copre piu' d'uno."""
    p = provenienze.get(chiave)
    for sotto in (p or {}).get("sottogruppi", []):
        if commento in sotto["voci"]:
            return sotto, chiave + "#" + sotto["nome"]
    return p, chiave


def origini():
    """Ogni record dei lotti, nella forma che il salvataggio contiene, con provenienza e dettaglio grezzo."""
    noti = {}
    for tipo, cartella in (("biglietto", "lotto-incontri-gen3"), ("scambio", "lotto-scambi-gen3")):
        for f in sorted(NOTE.joinpath(cartella).glob("*.pk3")):
            noti[NOTE.joinpath(cartella, "forma-cifrata", f.stem + ".ek3").read_bytes()] = (tipo, f.stem)
    for f in sorted(NOTE.joinpath("lotto-giganti").glob("*-gigante.bin")):
        noti[f.read_bytes()] = ("gigante", "86,2 cm, il massimo per la specie, da mostrare ai fratelli di Ceneride")
    manifesto = json.loads(NOTE.joinpath("lotto-wynaut", "manifesto.json").read_text(encoding="utf-8"))
    for f in sorted(NOTE.joinpath("lotto-wynaut").glob("*.bin")):
        voce = manifesto[f.stem.split("-")[0].capitalize()]
        noti[f.read_bytes()] = ("selvatico", "%s, livello %d, Ultra Ball: l'isola compare sul Percorso 130 e il gioco "
                                "registra quel percorso come luogo d'incontro" % (voce["luogo"], voce["livello"]))
    percorso = _modulo(CARTELLA.joinpath("tools", "parco_lotta_percorso_oro.py"), "percorso_oro")
    catalogo = json.loads(percorso.CATALOGO.read_text(encoding="utf-8"))
    thread = json.loads(percorso._mappa().THREAD.read_text(encoding="utf-8"))
    _, _, _, titolari = percorso.disposizione(catalogo, thread)
    for f in sorted(NOTE.joinpath("lotto-parco-lotta", "esemplari").glob("*-copia1.bin")):
        chiave = f.stem[:-len("-copia1")]
        noti[f.read_bytes()] = ("lotto", "titolare" if chiave in titolari else "riserva")
    return noti


def righe_dump(coppie):
    """Le righe dei dump per contenuto del record: {record: riga}, dalle posizioni dei file a cui i dump appartengono."""
    per_record = {}
    for dump, file_del_dump in coppie:
        salvataggio = save3.Save3(Path(file_del_dump).read_bytes())
        per_posizione = {}
        with open(dump, encoding="utf-8-sig") as f:
            for riga in csv.DictReader(f):
                m = re.search(r"@ \[(\d+)\] \([^)]*\)-(\d+)", riga["Position"])
                if m:
                    per_posizione[(int(m.group(1)) - 1) * PER_BOX + int(m.group(2)) - 1] = riga
        for i in range(save3.POSIZIONI):
            r = salvataggio.leggi_posizione(i)
            if r != VUOTO:
                if i not in per_posizione:
                    sys.exit("il dump %s non ha la posizione %d del suo file" % (dump, i))
                per_record.setdefault(r, per_posizione[i])
    return per_record


def riga_dai_byte(record, nature):
    """I campi della tabella per un record che nessun dump contiene, ricavati dai byte e dalle tabelle del progetto."""
    dati = json.loads(CARTELLA.joinpath("dati-gen3.json").read_text(encoding="utf-8"))
    mon = gen3.Gen3Mon.from_bytes(record)
    nome, info = next((n, v) for n, v in dati["specie"].items() if v["id"] == mon.growth.species)
    curva = dati["esperienza"][info["gruppo_crescita"]]
    livello = max(liv for liv in range(1, 101) if curva[liv] <= mon.growth.experience)
    tabella = charmap.Charmap.gen3()
    return {"Species": nome, "Nickname": "", "IsNicknamed": "False", "Level": str(livello),
            "OT": tabella.decode(mon.ot_name).strip(), "Nature": nature.get(mon.personality % 25, ""),
            "Form": "0", "MetLoc": "", "Version": "Smeraldo", "OTLang": "Italian",
            "Position": ": %04d" % nazionale_per_nome(nome)}


def nazionale_per_nome(nome):
    """Il numero nazionale dal nome inglese della specie, letto dai nomi dei file della raccolta di Sugimori: `dati-gen3.json` porta soltanto l'indice interno, che per le specie di Hoenn e' diverso."""
    for n in (1, 2, 3):
        for f in SUGIMORI.joinpath("Generation %d Pokémon" % n).glob("*.png"):
            m = re.match(r"(\d{4}) (.+)\.png$", f.name)
            if m and m.group(2).lower() == nome.lower():
                return int(m.group(1))
    sys.exit("specie %s assente dalla raccolta di Sugimori" % nome)


def illustrazioni():
    """Il file di Sugimori per numero nazionale, dalle prime tre generazioni."""
    indice = {}
    for n in (1, 2, 3):
        for f in SUGIMORI.joinpath("Generation %d Pokémon" % n).glob("*.png"):
            m = re.match(r"(\d{4}) ", f.name)
            if m:
                indice[int(m.group(1))] = f
    return indice, SUGIMORI.joinpath("Generation 3 Pokémon", "Pokémon Specific Forms Gen 3")


def miniatura(percorso, cache):
    if percorso not in cache:
        im = Image.open(percorso).convert("RGBA")
        # l'artwork di PKHeX e' piccolo, 68 per 56: si ingrandisce con un filtro morbido invece di lasciarlo sgranato
        if max(im.size) < 200:
            fattore = 220 / max(im.size)
            im = im.resize((round(im.width * fattore), round(im.height * fattore)), Image.LANCZOS)
        im.thumbnail((260, 260))
        cache[percorso] = im
    return cache[percorso]


def figura(titolo, voci, indice, forme, cache, uscita):
    """La griglia del box in proporzione A4 orizzontale, testo nero, senza legenda."""
    fig, ax = plt.subplots(figsize=(11.69, 8.27), facecolor=SUPERFICIE)
    fig.subplots_adjust(left=0.02, right=0.98, top=0.98, bottom=0.02)
    ax.set_xlim(0, COLONNE)
    ax.set_ylim(0, 5.3)
    ax.invert_yaxis()
    ax.set_aspect("auto")
    ax.axis("off")
    ax.text(0.04, 0.15, titolo, fontsize=17, color=NERO, fontweight="bold", va="center")
    for n, v in enumerate(voci, start=1):
        r, c = (n - 1) // COLONNE, (n - 1) % COLONNE
        y0 = r + 0.3
        ax.add_patch(FancyBboxPatch((c + 0.04, y0 + 0.04), 0.92, 0.92, boxstyle="round,pad=0,rounding_size=0.05",
                                    facecolor=SUPERFICIE, edgecolor=BORDO, linewidth=0.9))
        ax.text(c + 0.1, y0 + 0.13, str(n), fontsize=8, color=NERO, va="center")
        if not v:
            continue
        # L'etichetta porta il nome intero dell'evento: va a capo su due righe, e su tre con un corpo piu'
        # piccolo, invece di essere troncata con i puntini, che era il difetto della prima stampa.
        righe_banda = textwrap.wrap(v["banda"], 34) or [""]
        corpo_banda = 6.4 if len(righe_banda) <= 2 else 5.4
        ax.add_patch(FancyBboxPatch((c + 0.04, y0 + 0.76), 0.92, 0.20, boxstyle="round,pad=0,rounding_size=0.04",
                                    facecolor=TINTA[v["provenienza"]], edgecolor="none"))
        ax.text(c + 0.5, y0 + 0.86, "\n".join(righe_banda[:3]), fontsize=corpo_banda, color=NERO, ha="center", va="center",
                linespacing=1.05)
        percorso = indice.get(v["nazionale"])
        if v["nazionale"] == 201:
            f = int(v["forma"] or 0)
            percorso = ARTWORK_PKHEX.joinpath("a_201.png" if f == 0 else "a_201-%d.png" % f)
        if v["nazionale"] == 386 and v["forma"] in FORME_DEOXYS:
            alternativa = forme.joinpath("0386 Deoxys %s.png" % FORME_DEOXYS[v["forma"]])
            percorso = alternativa if alternativa.exists() else percorso
        if percorso:
            # la griglia riempie il foglio, quindi un'unita' orizzontale e una verticale non hanno la stessa
            # lunghezza: il lato orizzontale dell'immagine si corregge perche' l'illustrazione resti in proporzione
            lato = 0.42
            largo = lato * (8.27 * 0.96 / 5.3) / (11.69 * 0.96 / COLONNE)
            ax.imshow(miniatura(percorso, cache), extent=(c + 0.5 - largo / 2, c + 0.5 + largo / 2, y0 + 0.07 + lato, y0 + 0.07),
                      zorder=3, aspect="auto")
        ax.text(c + 0.5, y0 + 0.58, v["specie"], fontsize=9, color=NERO, ha="center", va="center", fontweight="bold")
        ax.text(c + 0.5, y0 + 0.68, ("Uovo   %s" if v["uovo"] else "Liv. %s   %%s" % v["livello"]) % v["allenatore"], fontsize=7.4, color=NERO,
                ha="center", va="center")
    fig.savefig(uscita, dpi=170, facecolor=SUPERFICIE, metadata={"Software": None})
    plt.close(fig)


def gruppo_rubino(motivo):
    """La provenienza di una richiesta del complemento, dal motivo che il manifesto le ha scritto."""
    for prefisso, chiave in (("statico", "r-statico"), ("incontro speciale", "r-statico"), ("esclusivo", "r-esclusivo"), ("specie", "r-specie"), ("mossa", "r-portatore"),
                             ("fiocchi", "r-portatore"), ("forma di Unown", "r-unown"), ("evento", "r-evento")):
        if motivo.startswith(prefisso):
            return chiave
    return "r-xd" if "XD" in motivo.split(",")[0] else "r-colosseum"


def disposizione_rubino():
    """Le voci del Rubino nell'ordine delle scatole: per gruppo, e dentro il gruppo per numero del Pokédex.

    E' la disposizione che la scrittura sulla cartuccia dovra' usare, e sta qui, in una funzione sola, perche'
    la mappa e la scrittura non possano divergere.
    """
    rapporto = json.loads(RUBINO.joinpath("rapporto.json").read_text(encoding="utf-8"))
    nomi = [n.strip() for n in PKHEX.joinpath("PKHeX.Core", "Resources", "text", "other", "it", "text_Species_it.txt").read_text(encoding="utf-8").split("\n")]
    ordine = [k for k, _, _, _ in PROVENIENZE_RUBINO]
    voci = []
    for e in rapporto["esiti"]:
        chiave = gruppo_rubino(e["motivo"])
        # la lettera di Unown si ricava dalla personalita' del file, come fa il gioco, e non dall'etichetta della
        # richiesta: il primo lotto aveva etichette da A a ? su 28 esemplari che erano tutti A
        forma = 0
        if chiave == "r-unown":
            pid = gen3.Gen3Mon.from_bytes(save3.record_da_file(RUBINO.joinpath(e["file"]).read_bytes())).personality
            forma = (((pid >> 24) & 3) << 6 | ((pid >> 16) & 3) << 4 | ((pid >> 8) & 3) << 2 | (pid & 3)) % 28
        dettaglio = e["motivo"].split(",", 1)[1].strip() if "," in e["motivo"] else e["motivo"]
        voci.append({"specie": nomi[e["specie"]] + (" " + LETTERE_UNOWN[forma] if chiave == "r-unown" else ""),
                     "nazionale": e["specie"], "forma": str(forma), "soprannome": "", "uovo": False,
                     "livello": str(e["livello"]), "allenatore": e["allenatore"].rsplit(" ", 1)[0], "lingua": "",
                     "provenienza": chiave, "banda": BANDA[chiave], "chiave": None, "file": e["file"],
                     "dettaglio": "%s; %s, allenatore %s" % (dettaglio, e["voce"], e["allenatore"])})
    voci.sort(key=lambda v: (ordine.index(v["provenienza"]), v["nazionale"], int(v["forma"])))
    return voci


def racconto(chiave, p, riga):
    """La storia dell'evento come testo continuo: nome, allenatore, lingua, quando, dove, come e perche' conta."""
    ot, ident = chiave.rsplit("|", 1)
    # per le uova l'identificativo della tabella e' zero, perche' e' quello di chi riceve: si scrive quello vero
    testa = "allenatore %s, identificativo %s, lingua %s" % (riga["OT"] or ot, riga.get("TID16") or ident, riga["OTLang"])
    if p is None:
        # ogni evento del deposito ha una voce in provenienze-eventi.json dal 2026-09-23: se ne manca una,
        # e' una lacuna del catalogo da colmare, non un testo da inventare qui
        sys.exit("l'evento %s non ha una voce in provenienze-eventi.json" % chiave)
    parti = ["%s: %s." % (p["nome"], testa), "Quando: %s." % p["date"], "Dove: %s." % p["luogo"], "Come: %s." % p["come"]]
    if p.get("oggetto_tenuto"):
        parti.append("Oggetto tenuto: %s." % p["oggetto_tenuto"].rstrip("."))
    if p.get("note"):
        parti.extend(frase for frase in re.split(r"(?<=[.;])\s+", p["note"]) if not DI_LAVORO.search(frase))
    return " ".join(parti)


def main():
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("salvataggio")
    p.add_argument("--dump", nargs=2, action="append", required=True, metavar=("DUMP", "FILE_DEL_DUMP"))
    args = p.parse_args()
    salvataggio = save3.Save3(Path(args.salvataggio).read_bytes())
    if not salvataggio.integro():
        sys.exit("lo slot attivo non e' integro")
    noti = origini()
    eventi, provenienze = storie_degli_eventi()
    per_record = righe_dump(args.dump)
    nature = {int(r["PID"], 16) % 25: r["Nature"] for r in per_record.values()}
    indice, forme = illustrazioni()

    voci = []
    specie_viste = set()
    for i in range(save3.POSIZIONI):
        r = salvataggio.leggi_posizione(i)
        if r == VUOTO:
            voci.append(None)
            continue
        riga = per_record.get(r)
        if riga is None:
            riga = riga_dai_byte(r, nature)
            print("posizione %d: nessun dump contiene questo record, descritto dai byte" % i)
        chiave, commento = eventi.get(r, (None, None))
        provenienza, dettaglio = ("evento", None) if chiave else noti.get(r, ("collezione", None))
        nazionale = int(re.search(r": (\d{4})\b", riga["Position"]).group(1))
        banda = BANDA[provenienza]
        if provenienza == "collezione":
            dettaglio = "%s, %s" % (riga["MetLoc"], riga["Version"])
            if nazionale in specie_viste:
                dettaglio += ", seconda cattura della specie"
            specie_viste.add(nazionale)
        elif provenienza == "evento":
            prov, identita = provenienza_evento(chiave, commento, provenienze)
            banda = prov["nome"] if prov else "Evento"
            dettaglio = racconto(chiave, prov, riga)
            chiave = identita
        elif provenienza == "biglietto":
            dettaglio = "%s, %s, catturato da un amico con il biglietto dell'evento e ricevuto in scambio" % (
                riga["MetLoc"], riga["Version"])
        elif provenienza == "scambio":
            dettaglio = "ricevuto in %s da %s, scambio a personalità fissa del gioco" % (riga["Version"], riga["OT"])
        elif provenienza == "lotto":
            banda = "Parco Lotta, %s" % dettaglio
            dettaglio = "%s, natura %s" % (dettaglio, riga["Nature"])
        voci.append({"specie": riga["Species"], "nazionale": nazionale, "forma": riga["Form"],
                     "soprannome": "uovo da schiudere" if riga.get("IsEgg") == "True" else
                     (riga["Nickname"] if riga["IsNicknamed"] == "True" else ""),
                     "uovo": riga.get("IsEgg") == "True",
                     "livello": riga["Level"], "allenatore": ("<AMICO>" if provenienza == "collezione" and riga["OT"] not in OT_COLLEZIONE_PUBBLICI else riga["OT"]), "lingua": riga["OTLang"],
                     "provenienza": provenienza, "dettaglio": dettaglio, "banda": banda, "chiave": chiave})

    # la storia di un evento si scrive alla prima posizione del box che lo contiene, le altre vi rimandano
    per_box = [voci[b * PER_BOX:(b + 1) * PER_BOX] for b in range(save3.POSIZIONI // PER_BOX)]
    for contenuto in per_box:
        prima = {}
        for n, v in enumerate(contenuto, start=1):
            if v and v["chiave"]:
                if v["chiave"] in prima:
                    v["dettaglio"] = "Stesso evento della posizione %d; allenatore %s, lingua %s." % (
                        prima[v["chiave"]], v["allenatore"], v["lingua"])
                else:
                    prima[v["chiave"]] = n

    FIGURE.mkdir(exist_ok=True)
    cache = {}
    for b, contenuto in enumerate(per_box, start=1):
        figura("Smeraldo, box %d" % b, contenuto, indice, forme, cache, FIGURE.joinpath("mappa-box-%02d.png" % b))

    completa = _modulo(CARTELLA.joinpath("tools", "emerald_cartuccia_completa.py"), "cartuccia_completa")
    sfondi = list(salvataggio.storage()[save3.OFF_SFONDI:save3.OFF_SFONDI + 14])
    conteggio = collections.Counter(v["provenienza"] for v in voci if v)
    riepilogo = []
    for b, contenuto in enumerate(per_box, start=1):
        parti = collections.Counter(v["provenienza"] for v in contenuto if v)
        riepilogo.append((b, "; ".join("%s %d" % (ETICHETTA[k], parti[k]) for k, _, _, _ in PROVENIENZE if parti[k]),
                          completa.NOMI_SFONDI[sfondi[b - 1]]))
    scrivi_markdown(per_box, conteggio, riepilogo, PROVENIENZE, USCITA, None)
    sezioni = [{"nome": "Smeraldo", "per_box": per_box, "conteggio": conteggio, "riepilogo": riepilogo, "provenienze": PROVENIENZE,
                "figura": "mappa-box-%02d.png", "nota": None}]

    # il Rubino: disposizione prevista del complemento, finche' la cartuccia non e' scritta
    voci_r = disposizione_rubino()
    per_box_r = [voci_r[b * PER_BOX:(b + 1) * PER_BOX] + [None] * max(0, PER_BOX - len(voci_r[b * PER_BOX:(b + 1) * PER_BOX]))
                 for b in range(-(-len(voci_r) // PER_BOX))]
    for b, contenuto in enumerate(per_box_r, start=1):
        figura("Rubino, box %d (previsto)" % b, contenuto, indice, forme, cache, FIGURE.joinpath("mappa-rubino-box-%02d.png" % b))
    conteggio_r = collections.Counter(v["provenienza"] for v in voci_r)
    sfondi_r = sfondi_rubino(per_box_r)
    riepilogo_r = [(b, "; ".join("%s %d" % (ETICHETTA[k], collections.Counter(v["provenienza"] for v in c if v)[k])
                                 for k, _, _, _ in PROVENIENZE_RUBINO if collections.Counter(v["provenienza"] for v in c if v)[k]),
                    NOMI_SFONDI_RUBINO[sfondi_r[b - 1]])
                   for b, c in enumerate(per_box_r, start=1)]
    nota_r = ("Disposizione prevista del complemento di ADR-080, non ancora scritta: il Rubino di prova non è a portata di mano. "
              "La squadra di inizio partita resta quella che la cartuccia ha, e le scatole oltre queste restano vuote.")
    scrivi_markdown(per_box_r, conteggio_r, riepilogo_r, PROVENIENZE_RUBINO, USCITA_RUBINO, nota_r)
    sezioni.append({"nome": "Rubino", "per_box": per_box_r, "conteggio": conteggio_r, "riepilogo": riepilogo_r,
                    "provenienze": PROVENIENZE_RUBINO, "figura": "mappa-rubino-box-%02d.png", "nota": nota_r})
    scrivi_stampa(sezioni)
    for k, e, _, _ in PROVENIENZE + PROVENIENZE_RUBINO:
        print("  %-38s %d" % (e, (conteggio + conteggio_r)[k]))


def scrivi_markdown(per_box, conteggio, riepilogo, provenienze, uscita, nota):
    md = ["<!-- generato da gba-save-extraction-smeraldo/tools/emerald_mappa_box.py: non si modifica a mano -->", ""]
    if uscita == USCITA:
        md += ["`<AMICO>` maschera il nome di un allenatore terzo nella mappa pubblica; i record del salvataggio non sono cambiati.", ""]
    if nota:
        md += [nota, ""]
    md += ["## Legenda e riepilogo", "", "| Colore | Provenienza | Esemplari |", "|---|---|---|"]
    for k, e, _, _ in provenienze:
        md.append("| %s | %s | %d |" % (TINTA[k], e, conteggio[k]))
    md.append("|  | Totale | %d su %d |" % (sum(conteggio.values()), save3.POSIZIONI))
    md += ["", "| Box | Contenuto | Sfondo |", "|---|---|---|"]
    md += ["| %d | %s | %s |" % t for t in riepilogo]
    for b, contenuto in enumerate(per_box, start=1):
        md += ["", "## Box %d" % b, "", "![Box %d](figure/%s)" % (b, ("mappa-rubino-box-%02d.png" if uscita == USCITA_RUBINO else "mappa-box-%02d.png") % b), "",
               "| Posizione | Riga | Colonna | Pokémon | Soprannome | Livello | Allenatore | Provenienza | Dettaglio |",
               "|---|---|---|---|---|---|---|---|---|"]
        for n, v in enumerate(contenuto, start=1):
            r, c = (n - 1) // COLONNE + 1, (n - 1) % COLONNE + 1
            if not v:
                md.append("| %d | %d | %d | vuoto |  |  |  |  |  |" % (n, r, c))
                continue
            md.append("| %d | %d | %d | %s | %s | %s | %s | %s | %s |" % (
                n, r, c, v["specie"], v["soprannome"], v["livello"], v["allenatore"], ETICHETTA[v["provenienza"]],
                v["dettaglio"].replace("|", "/")))
    uscita.write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")


def scrivi_stampa(sezioni):
    """La copia da stampare: .docx orizzontale costruito qui, esportato in PDF da Word, con il controllo delle pagine."""
    corpi = {(s["nome"], b): 7.0 for s in sezioni for b in range(1, len(s["per_box"]) + 1)}
    passo = 0.25
    docx_out, pdf_out = STAMPA.with_suffix(".docx"), STAMPA.with_suffix(".pdf")
    for tentativo in range(8):
        componi_docx(sezioni, corpi, docx_out)
        esporta_pdf(docx_out, pdf_out)
        sforano, pagine = pagine_che_sforano(pdf_out, sezioni)
        if not sforano:
            print("stampa: %s e %s, %d pagine: per cartuccia una di legenda e due per box" % (docx_out.name, pdf_out.name, pagine))
            return
        for b in sforano:
            corpi[b] -= passo
            if corpi[b] < 4.5:
                sys.exit("la tabella del box %s non entra in una pagina neppure a 4,5 punti" % (b,))
        print("tentativo %d: sforano i box %s, riduco il carattere" % (tentativo + 1, sforano))
    sys.exit("impaginazione non riuscita")


def componi_docx(sezioni, corpi, uscita):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sezione = doc.sections[0]
    sezione.orientation = WD_ORIENT.LANDSCAPE
    sezione.page_width, sezione.page_height = Cm(29.7), Cm(21.0)
    for lato in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sezione, lato, Cm(0.9))
    stile = doc.styles["Normal"]
    stile.font.name = "Arial"
    stile.font.size = Pt(2)
    stile.font.color.rgb = RGBColor(0, 0, 0)
    stile.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "MS Gothic")
    stile.paragraph_format.space_after = Pt(0)
    stile.paragraph_format.space_before = Pt(0)

    def testo(paragrafo, contenuto, corpo, grassetto=False, font="Arial Narrow"):
        run = paragrafo.add_run(contenuto)
        run.font.size = Pt(corpo)
        run.font.name = font
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "MS Gothic")
        run.font.bold = grassetto
        run.font.color.rgb = RGBColor(0, 0, 0)
        return run

    def ombra(cella, colore):
        tcPr = cella._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), colore.lstrip("#"))
        tcPr.append(shd)

    def margini_stretti(t, cm=0.06):
        mar = OxmlElement("w:tblCellMar")
        for lato in ("top", "bottom", "left", "right"):
            e = OxmlElement("w:%s" % lato)
            e.set(qn("w:w"), str(int(cm * 567)))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        t._element.tblPr.append(mar)

    def tabella(intestazione, righe, larghezze, corpo, colori=None):
        t = doc.add_table(rows=1 + len(righe), cols=len(intestazione))
        t.style = "Table Grid"
        t.autofit = False
        margini_stretti(t)
        for j, (h, w) in enumerate(zip(intestazione, larghezze)):
            cella = t.rows[0].cells[j]
            cella.width = Cm(w)
            testo(cella.paragraphs[0], h, corpo, grassetto=True)
        for i, riga in enumerate(righe, start=1):
            for j, (valore, w) in enumerate(zip(riga, larghezze)):
                cella = t.rows[i].cells[j]
                cella.width = Cm(w)
                testo(cella.paragraphs[0], valore, corpo)
            if colori and colori[i - 1]:
                ombra(t.rows[i].cells[colori[i - 1][0]], colori[i - 1][1])
        return t

    def a_capo_pagina():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for i, sezione in enumerate(sezioni):
        if i:
            a_capo_pagina()
        conteggio, provenienze = sezione["conteggio"], sezione["provenienze"]
        testo(doc.add_paragraph(), "%s: legenda e riepilogo" % sezione["nome"], 12, grassetto=True, font="Arial")
        if sezione["nota"]:
            testo(doc.add_paragraph(), sezione["nota"], 9, font="Arial")
        testo(doc.add_paragraph(), " ", 8)
        tabella(["Colore", "Provenienza", "Esemplari"],
                [("", e, str(conteggio[k])) for k, e, _, _ in provenienze] + [("", "Totale", "%d su %d" % (sum(conteggio.values()), save3.POSIZIONI))],
                [1.5, 8.0, 3.0], 9, colori=[(0, TINTA[k]) for k, _, _, _ in provenienze] + [None])
        testo(doc.add_paragraph(), " ", 8)
        tabella(["Box", "Contenuto", "Sfondo"], [(str(b), c, sf) for b, c, sf in sezione["riepilogo"]], [1.5, 20.0, 3.0], 9)
        for b, contenuto in enumerate(sezione["per_box"], start=1):
            a_capo_pagina()
            doc.add_paragraph().add_run().add_picture(str(FIGURE.joinpath(sezione["figura"] % b)), height=Cm(18.4))
            a_capo_pagina()
            testo(doc.add_paragraph(), "%s, box %d, posizione per posizione" % (sezione["nome"], b), 10, grassetto=True, font="Arial")
            righe, colori = [], []
            for n, v in enumerate(contenuto, start=1):
                rc = "%d (%d-%d)" % (n, (n - 1) // COLONNE + 1, (n - 1) % COLONNE + 1)
                if not v:
                    righe.append((rc, "vuoto", "", "", "", "", ""))
                    colori.append(None)
                    continue
                righe.append((rc, v["specie"], v["soprannome"], "uovo" if v["uovo"] else v["livello"], v["allenatore"], ETICHETTA[v["provenienza"]], v["dettaglio"]))
                colori.append((5, TINTA[v["provenienza"]]))
            tabella(["Posto (riga-colonna)", "Pokémon", "Soprannome", "Liv.", "Allenatore", "Provenienza", "Dettaglio"],
                    righe, [1.4, 1.7, 1.5, 0.6, 1.7, 1.9, 19.1], corpi[(sezione["nome"], b)], colori)
    doc.save(str(uscita))


def esporta_pdf(docx_in, pdf_out):
    """Word esporta il .docx in PDF: e' l'impaginazione che il proprietario stampera', non una stima."""
    comando = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
               "$d = $w.Documents.Open('%s'); $d.ExportAsFixedFormat('%s', 17); $d.Close($false); $w.Quit()"
               % (str(docx_in.resolve()), str(pdf_out.resolve())))
    # la prima apertura di Word dopo un avvio a volte esce senza scrivere: si cancella il PDF vecchio,
    # cosi' che non passi per nuovo, e si riprova una volta prima di arrendersi
    for _ in range(2):
        if pdf_out.exists():
            pdf_out.unlink()
        subprocess.run(["powershell", "-NoProfile", "-Command", comando], check=True)
        if pdf_out.exists():
            return
    sys.exit("Word non ha esportato %s" % pdf_out)


def pagine_che_sforano(pdf, sezioni):
    """I box la cui tabella occupa piu' di una pagina: dopo la pagina con il titolo della tabella deve venire la figura del box seguente."""
    import fitz
    documento = fitz.open(str(pdf))
    titoli = []
    for i, pagina in enumerate(documento):
        m = re.search(r"(\w+), box (\d+), posizione per posizione", pagina.get_text())
        if m:
            titoli.append(((m.group(1), int(m.group(2))), i))
    # la pagina di legenda della cartuccia seguente sta fra l'ultima tabella di una cartuccia e la figura del primo box
    fine_sezione = {s["nome"]: len(s["per_box"]) for s in sezioni}
    sforano = []
    for j, ((nome, b), inizio) in enumerate(titoli):
        dopo = titoli[j + 1][1] if j + 1 < len(titoli) else len(documento) + 1
        attese = 2 if b < fine_sezione[nome] else 3
        if j + 1 == len(titoli):
            attese = 1
            dopo = len(documento)
        if dopo - inizio > attese:
            sforano.append((nome, b))
    pagine = len(documento)
    documento.close()
    return sforano, pagine


if __name__ == "__main__":
    main()
